"""
Builds the final Affidavit in Reply .docx from:
  - `context` (fixed facts, from case_info_parser.to_affidavit_context)
  - `drafted` (prose, from draft_engine.draft_affidavit_prose)

All layout/formatting here is deterministic - the same context+drafted
shape always produces the same structure, regardless of what the LLM
happened to phrase things as. This is what makes the pipeline reliable:
the model only supplies sentences, never document structure.
"""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re


def _ordinal_date(date_str):
    """'5 September 2026' -> '5th day of September 2026'. Falls back to the
    raw string unchanged if it doesn't match the expected 'D Month YYYY' shape."""
    m = re.match(r'^\s*(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})\s*$', date_str or '')
    if not m:
        return date_str
    day, month, year = m.groups()
    day_i = int(day)
    suffix = 'th' if 11 <= day_i % 100 <= 13 else {1: 'st', 2: 'nd', 3: 'rd'}.get(day_i % 10, 'th')
    return f"{day_i}{suffix} day of {month} {year}"


def _no_border(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'none')
        el.set(qn('w:sz'), '0')
        el.set(qn('w:space'), '0')
        borders.append(el)
    tblPr.append(borders)


def _two_col_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    _no_border(table)
    for left, right in rows:
        row = table.add_row()
        row.cells[0].paragraphs[0].add_run(left)
        p = row.cells[1].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run(right)
    return table


def _centered(doc, text, bold=True, size=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    return p


def _justified(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(text)
    return p


def _numbered(doc, marker, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Pt(28)
    p.paragraph_format.first_line_indent = Pt(-28)
    p.add_run(f"{marker}.\t{text}")
    return p


def build_affidavit(context, drafted, output_path):
    case = context['case']
    deponent = context['deponent']
    attestation = context['attestation']
    advocate = context['advocate']

    doc = Document()

    _centered(doc, case.get('Court', ''))
    if case.get('Jurisdiction'):
        _centered(doc, case['Jurisdiction'])
    proceeding = case.get('Proceeding Type', 'PETITION')
    _centered(
        doc,
        f"{proceeding} NO. {case.get('Case Number', '')} OF {case.get('Year', '')}",
        size=13,
    )
    doc.add_paragraph()

    _two_col_table(doc, [(case.get('Petitioner', ''), ''), ('', '...Petitioner')])
    doc.add_paragraph()
    _centered(doc, 'VERSUS')
    doc.add_paragraph()

    respondent_rows = []
    if case.get('Respondent No. 1'):
        respondent_rows.append((f"1. {case['Respondent No. 1']}", ''))
        respondent_rows.append(('', '...Respondent No.1'))
    if case.get('Respondent No. 2'):
        respondent_rows.append((f"2. {case['Respondent No. 2']}", ''))
        respondent_rows.append(('', '...Respondent No.2'))
    _two_col_table(doc, respondent_rows)

    filed_for = case.get('Filed on behalf of', '')
    doc_type = case.get('Document Type', 'AFFIDAVIT')
    _centered(doc, f"{doc_type.upper()} ON BEHALF OF {filed_for.upper()}")

    _justified(doc, drafted['opening'])

    for i, para_text in enumerate(drafted['paragraphs'], start=1):
        _numbered(doc, str(i), para_text)

    _centered(doc, 'PRAYER')
    _justified(doc, drafted.get('prayer_intro', 'I therefore respectfully pray that this Hon\u2019ble Court may be pleased to:'))
    letters = 'abcdefghijklmnopqrstuvwxyz'
    for i, item in enumerate(drafted['prayer_items']):
        suffix = ';' if i < len(drafted['prayer_items']) - 1 else '.'
        _numbered(doc, f"({letters[i]})", f"{item}{suffix}")

    formatted_date = _ordinal_date(attestation.get('Date', ''))

    doc.add_paragraph()
    _two_col_table(doc, [
        (f"Solemnly affirmed at {attestation.get('Place', '')}", ''),
        (f"On this {formatted_date}", ''),
        ('Before Me', 'DEPONENT'),
    ])

    _centered(doc, 'VERIFICATION')
    n_paras = len(drafted['paragraphs'])
    _justified(
        doc,
        f"I, {deponent.get('Name', '')}, the Deponent above named, do hereby verify "
        f"that the contents of paragraphs 1 to {n_paras} and the Prayer above are true "
        "and correct to my knowledge and belief and that nothing material has been "
        "concealed therefrom."
    )
    _justified(doc, f"Verified at {attestation.get('Place', '')} on this {formatted_date}.")

    _centered(doc, 'DEPONENT')

    doc.add_paragraph()
    _centered(doc, advocate.get('Advocate Firm', '').upper())
    _centered(doc, f"Advocates for the {advocate.get('Acting for', '')}.", bold=False)

    doc.save(output_path)
    return output_path


if __name__ == '__main__':
    import sys, json
    context_path, drafted_path, output_path = sys.argv[1], sys.argv[2], sys.argv[3]
    with open(context_path) as f:
        context = json.load(f)
    with open(drafted_path) as f:
        drafted = json.load(f)
    build_affidavit(context, drafted, output_path)
    print(f"Wrote {output_path}")
